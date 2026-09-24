from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger(__name__)

EURO_FORMAT = "#,##0.00 [$€-it-IT]"
HEADER_FILL = PatternFill("solid", fgColor="FFBDD7EE")
HEADER_FONT = Font(bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FF00B050")
TITLE_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
TOTAL_LABEL = "Total"


def export_report(dataframe: pd.DataFrame, pivot: pd.DataFrame, output_path: Path) -> None:
    """Write raw data and pivot sheets, then apply professional formatting."""
    
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name="DATI", index=False)
        pivot.to_excel(writer, sheet_name="REPORT", index=False, startrow=2)

    workbook = load_workbook(output_path)
    _format_data_sheet(workbook["DATI"], dataframe)
    _format_report_sheet(workbook["REPORT"], pivot)
    workbook.save(output_path)
    LOGGER.info("Report Excel salvato in %s", output_path)


def _format_data_sheet(worksheet, dataframe: pd.DataFrame) -> None:
    _style_header_row(worksheet, header_row=1)
    _apply_borders(worksheet)
    _auto_fit_columns(worksheet)
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"

    amount_column = _find_column_letter(worksheet, "IMPORTO", header_row=1)
    if amount_column:
        for row_index in range(2, len(dataframe) + 2):
            worksheet[f"{amount_column}{row_index}"].number_format = EURO_FORMAT


def _format_report_sheet(worksheet, pivot: pd.DataFrame) -> None:
    max_column = worksheet.max_column
    title_end = get_column_letter(max_column)
    worksheet.merge_cells(f"A1:{title_end}1")
    title_cell = worksheet["A1"]
    title_cell.value = "Report Vendite per Regione e Categoria"
    title_cell.fill = TITLE_FILL
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    _style_header_row(worksheet, header_row=3)
    _apply_borders(worksheet)
    _auto_fit_columns(worksheet)

    worksheet.auto_filter.ref = f"A3:{get_column_letter(max_column)}{worksheet.max_row}"
    worksheet.freeze_panes = "B4"

    total_row_index = None
    for row_index in range(4, worksheet.max_row + 1):
        label_value = worksheet[f"A{row_index}"].value
        if label_value == TOTAL_LABEL:
            total_row_index = row_index
            break

    for row_index in range(4, worksheet.max_row + 1):
        # Additional formatting logic can be added here
        pass


def _style_header_row(worksheet, header_row: int) -> None:
    for cell in worksheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _apply_borders(worksheet) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER


def _auto_fit_columns(worksheet) -> None:
    for column_index in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(column_index)
        max_length = 0
        for row_index in range(1, worksheet.max_row + 1):
            cell_value = worksheet.cell(row=row_index, column=column_index).value
            max_length = max(max_length, len(str(cell_value or "")))
        worksheet.column_dimensions[column_letter].width = max(max_length + 2, 12)


def _find_column_letter(worksheet, column_name: str, header_row: int) -> str | None:
    for column_index in range(1, worksheet.max_column + 1):
        if worksheet.cell(row=header_row, column=column_index).value == column_name:
            return get_column_letter(column_index)
    return None