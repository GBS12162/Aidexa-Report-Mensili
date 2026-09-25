"""Data extraction logic for the Oracle onboarding error report."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
import oracledb
from openpyxl import load_workbook

from config import get_config

LOGGER = logging.getLogger(__name__)

# Column layout of examples/caso1/query.sql (last two columns have no alias).
COLUMNS = get_config().query_columns

# Header labels expected on row 1 of examples/*/input.xlsx (order-independent lookup).
_INPUT_HEADER_NAMES = get_config().input_header_names


def load_query_text(query_file: Path) -> str:
    """Read the SQL query text from disk."""

    if not query_file.is_file():
        raise FileNotFoundError(f"File query non trovato: {query_file}")

    text = query_file.read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError(f"Il file query è vuoto: {query_file}")
    if text.endswith(";"):
        text = text[:-1].rstrip()
    return text


def extract_report_data(connection: oracledb.Connection, query_file: Path) -> pd.DataFrame:
    """Execute the fixed SQL query and return the full result set as a DataFrame."""

    query_text = load_query_text(query_file)

    with connection.cursor() as cursor:
        cursor.execute(query_text)
        rows = cursor.fetchall()

    dataframe = pd.DataFrame(rows, columns=COLUMNS)
    LOGGER.info("Estratte %s righe da Oracle", len(dataframe))
    return dataframe


def load_mock_dataframe_from_excel(input_file: Path) -> pd.DataFrame:
    """Load a DataFrame with the same shape as extract_report_data(), but read
    from an Excel dump of a query result set (used by the --test-excel mode
    to simulate the Oracle extraction without any database access)."""

    if not input_file.is_file():
        raise FileNotFoundError(f"File di input non trovato: {input_file}")

    workbook = load_workbook(input_file, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    header_row = [worksheet.cell(row=1, column=col).value for col in range(1, worksheet.max_column + 1)]
    column_positions = {}
    for name in _INPUT_HEADER_NAMES:
        try:
            column_positions[name] = header_row.index(name) + 1
        except ValueError as exc:
            raise ValueError(f"Colonna '{name}' non trovata in {input_file}") from exc

    # The value column ("NVL(...)") has no stable alias: it is positioned
    # right after TIPO_ERRORE, followed by the unused raw conteggio column.
    value_col = column_positions["TIPO_ERRORE"] + 1

    rows = []
    for row_index in range(2, worksheet.max_row + 1):
        url = worksheet.cell(row=row_index, column=column_positions["URLL"]).value
        if url is None:
            continue
        rows.append(
            (
                url,
                worksheet.cell(row=row_index, column=column_positions["DATAA"]).value,
                worksheet.cell(row=row_index, column=column_positions["STATOO"]).value,
                worksheet.cell(row=row_index, column=column_positions["TIPO_ERRORE"]).value,
                worksheet.cell(row=row_index, column=value_col).value or 0,
                None,
            )
        )

    dataframe = pd.DataFrame(rows, columns=COLUMNS)
    LOGGER.info("Caricate %s righe da %s (modalità test, nessuna connessione Oracle)", len(dataframe), input_file)
    return dataframe

