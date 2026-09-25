"""Structural and formatting validation of a generated report against a
reference workbook. Used by the --test-excel mode."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

MAX_REPORTED_PER_KIND = 30


@dataclass
class Difference:
    """A single detected difference between generated and reference workbook."""

    sheet: str
    kind: str
    coordinate: str | None
    expected: object
    actual: object
    row: int | None = None
    column: int | None = None

    def render(self) -> str:
        location = f"{self.coordinate}" if self.coordinate else "-"
        return (
            f"[{self.sheet}] tipo={self.kind} cella={location} "
            f"riga={self.row} colonna={self.column} "
            f"atteso={self.expected!r} generato={self.actual!r}"
        )


@dataclass
class ValidationReport:
    """Collects differences found while comparing two workbooks."""

    differences: list[Difference] = field(default_factory=list)
    _kind_counts: dict = field(default_factory=dict)

    @property
    def is_valid(self) -> bool:
        return not self.differences

    def add(self, sheet: str, kind: str, expected, actual, coordinate: str | None = None,
            row: int | None = None, column: int | None = None) -> None:
        count = self._kind_counts.get(kind, 0)
        self._kind_counts[kind] = count + 1
        if count < MAX_REPORTED_PER_KIND:
            self.differences.append(
                Difference(sheet=sheet, kind=kind, coordinate=coordinate, expected=expected,
                           actual=actual, row=row, column=column)
            )

    def render(self) -> str:
        if self.is_valid:
            return "Nessuna differenza rilevata. output_test.xlsx è identico al riferimento."

        lines = [f"Rilevate differenze per {len(self._kind_counts)} categorie:"]
        for kind, total in self._kind_counts.items():
            omitted = total - min(total, MAX_REPORTED_PER_KIND)
            lines.append(f"-- {kind}: {total} differenze" + (f" ({omitted} omesse)" if omitted else ""))
        lines.append("")
        lines.extend(diff.render() for diff in self.differences)
        return "\n".join(lines)


def validate_against_reference(generated_path: Path, reference_path: Path) -> ValidationReport:
    """Compare structure, styles and values between two xlsx files."""

    report = ValidationReport()

    if not generated_path.is_file():
        report.add("-", "file_mancante", str(reference_path), None)
        return report
    if not reference_path.is_file():
        report.add("-", "file_mancante", None, str(generated_path))
        return report

    generated = load_workbook(generated_path)
    reference = load_workbook(reference_path)

    if generated.sheetnames != reference.sheetnames:
        report.add("-", "fogli", reference.sheetnames, generated.sheetnames)

    for sheet_name in reference.sheetnames:
        if sheet_name not in generated.sheetnames:
            continue
        _compare_sheet(generated[sheet_name], reference[sheet_name], sheet_name, report)

    return report


def _compare_sheet(generated_ws: Worksheet, reference_ws: Worksheet, sheet_name: str, report: ValidationReport) -> None:
    if generated_ws.max_row != reference_ws.max_row:
        report.add(sheet_name, "numero_righe", reference_ws.max_row, generated_ws.max_row)
    if generated_ws.max_column != reference_ws.max_column:
        report.add(sheet_name, "numero_colonne", reference_ws.max_column, generated_ws.max_column)

    _compare_merged_cells(generated_ws, reference_ws, sheet_name, report)
    _compare_column_widths(generated_ws, reference_ws, sheet_name, report)
    _compare_row_heights(generated_ws, reference_ws, sheet_name, report)
    _compare_sheet_view(generated_ws, reference_ws, sheet_name, report)

    max_row = min(generated_ws.max_row, reference_ws.max_row)
    max_col = min(generated_ws.max_column, reference_ws.max_column)

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            generated_cell = generated_ws.cell(row=row, column=col)
            reference_cell = reference_ws.cell(row=row, column=col)
            _compare_cell(generated_cell, reference_cell, sheet_name, row, col, report)


def _compare_cell(generated_cell: Cell, reference_cell: Cell, sheet_name: str, row: int, col: int,
                   report: ValidationReport) -> None:
    coordinate = reference_cell.coordinate

    if generated_cell.value != reference_cell.value:
        report.add(sheet_name, "valore", reference_cell.value, generated_cell.value, coordinate, row, col)

    if generated_cell.number_format != reference_cell.number_format:
        report.add(sheet_name, "formato_numerico", reference_cell.number_format, generated_cell.number_format,
                    coordinate, row, col)

    ref_align = (reference_cell.alignment.horizontal, reference_cell.alignment.vertical)
    gen_align = (generated_cell.alignment.horizontal, generated_cell.alignment.vertical)
    if ref_align != gen_align and any(ref_align):
        report.add(sheet_name, "allineamento", ref_align, gen_align, coordinate, row, col)

    ref_font = _font_signature(reference_cell)
    gen_font = _font_signature(generated_cell)
    if ref_font != gen_font:
        report.add(sheet_name, "font", ref_font, gen_font, coordinate, row, col)

    ref_fill = _fill_signature(reference_cell)
    gen_fill = _fill_signature(generated_cell)
    if ref_fill != gen_fill:
        report.add(sheet_name, "colore_sfondo", ref_fill, gen_fill, coordinate, row, col)

    ref_border = _border_signature(reference_cell)
    gen_border = _border_signature(generated_cell)
    if ref_border != gen_border:
        report.add(sheet_name, "bordi", ref_border, gen_border, coordinate, row, col)


def _color_signature(color) -> tuple | None:
    if color is None:
        return None
    if color.type == "rgb":
        return ("rgb", color.rgb)
    if color.type == "theme":
        return ("theme", color.theme, round(color.tint or 0, 4))
    if color.type == "indexed":
        return ("indexed", color.indexed)
    return None


def _font_signature(cell: Cell) -> tuple:
    font = cell.font
    return (font.name, font.size, font.bold, font.italic, _color_signature(font.color))


def _fill_signature(cell: Cell) -> tuple | None:
    fill = cell.fill
    if fill.fill_type is None:
        return None
    return (fill.fill_type, _color_signature(fill.fgColor))


def _border_signature(cell: Cell) -> tuple:
    border = cell.border
    return tuple(
        getattr(border, side).style if getattr(border, side) else None
        for side in ("left", "right", "top", "bottom")
    )


def _compare_merged_cells(generated_ws: Worksheet, reference_ws: Worksheet, sheet_name: str,
                           report: ValidationReport) -> None:
    generated_ranges = {str(r) for r in generated_ws.merged_cells.ranges}
    reference_ranges = {str(r) for r in reference_ws.merged_cells.ranges}
    if generated_ranges != reference_ranges:
        missing = sorted(reference_ranges - generated_ranges)
        extra = sorted(generated_ranges - reference_ranges)
        if missing:
            report.add(sheet_name, "celle_unite_mancanti", missing, None)
        if extra:
            report.add(sheet_name, "celle_unite_extra", None, extra)


def _compare_column_widths(generated_ws: Worksheet, reference_ws: Worksheet, sheet_name: str,
                            report: ValidationReport) -> None:
    reference_widths = {key: dim.width for key, dim in reference_ws.column_dimensions.items() if dim.width}
    generated_widths = {key: dim.width for key, dim in generated_ws.column_dimensions.items() if dim.width}
    for column_letter, expected_width in reference_widths.items():
        actual_width = generated_widths.get(column_letter)
        if actual_width is None or round(actual_width, 2) != round(expected_width, 2):
            report.add(sheet_name, "larghezza_colonna", expected_width, actual_width, column_letter)


def _compare_row_heights(generated_ws: Worksheet, reference_ws: Worksheet, sheet_name: str,
                          report: ValidationReport) -> None:
    reference_heights = {key: dim.height for key, dim in reference_ws.row_dimensions.items() if dim.height}
    generated_heights = {key: dim.height for key, dim in generated_ws.row_dimensions.items() if dim.height}
    for row_index, expected_height in reference_heights.items():
        actual_height = generated_heights.get(row_index)
        if actual_height is None or round(actual_height, 2) != round(expected_height, 2):
            report.add(sheet_name, "altezza_riga", expected_height, actual_height, row=row_index)


def _compare_sheet_view(generated_ws: Worksheet, reference_ws: Worksheet, sheet_name: str,
                         report: ValidationReport) -> None:
    if generated_ws.freeze_panes != reference_ws.freeze_panes:
        report.add(sheet_name, "freeze_panes", reference_ws.freeze_panes, generated_ws.freeze_panes)

    reference_filter = reference_ws.auto_filter.ref
    generated_filter = generated_ws.auto_filter.ref
    if generated_filter != reference_filter:
        report.add(sheet_name, "filtri", reference_filter, generated_filter)

